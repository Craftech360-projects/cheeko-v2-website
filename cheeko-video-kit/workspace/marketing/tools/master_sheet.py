"""Cheeko Video Master Sheet: one spreadsheet in the Drive folder that lists every video with direct links
(watch, 16:9 cut, thumbnail, Ready to post, all files, script) and its caption, for the team.

Built from the same VIDEOS list as publish_to_drive.py, which calls build() at the end of every full run.
Links are Google Drive links read from Drive for desktop (xattr com.google.drivefs.item-id#S), so they open for
anyone the "cheeko ai videos" folder is shared with.
usage: master_sheet.py
"""
import datetime, importlib.util, subprocess
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
_spec = importlib.util.spec_from_file_location("publish", HERE / "publish_to_drive.py")
P = importlib.util.module_from_spec(_spec); _spec.loader.exec_module(P)
OUT_NAME = "Cheeko Video Master Sheet.xlsx"

ABOUT = {   # one line per video for the team (Hindi rows reuse the English line)
    "V01": "All 8 features in 22 seconds: cards, Talk, Imagine, games, languages, Radio, Funny Voice, parent app.",
    "V02": "The cards: \"Kid bored? Skip the phone. Hand them Cheeko.\"",
    "V03": "\"Took the phone away, and no tantrum?!\" A fast feature tour.",
    "V04": "\"Your phone is built to never end. Cheeko is built to END.\" The screen that stops.",
    "V05": "\"Guess this sound!\" Pressure cooker hook, Sounds Around Me card, Nani, Imagine, languages, made in India.",
    "V06": "Sticker ad: the family phone tells how \"five minutes\" became dinner, football and one more, then Cheeko showed up.",
    "V07": "Illustrated family moments: little moments, more connection.",
    "V08": "Sarcastic parent reel: a child inherits an adult's feed.",
    "V09": "Sticker ad: phone-raised kids at a 7th birthday act like tiny adults, then rewind to real play with Cheeko.",
    "V11": "Kids ask anything, Cheeko answers.",
    "V12": "Say a wish, Cheeko draws it.",
    "V13": "Record your voice, hear it back as chipmunk, monster, robot and more.",
    "V14": "8 games that work without Wi-Fi. No ads, no purchases.",
    "V15": "Grandma records clips on a custom card, the child plays them any time.",
    "V16": "Nani tells stories and listens when the child interrupts.",
    "V17": "Quizzy: 10 questions a day, answered in the child's own words.",
    "V18": "Mitthu, the English teacher: new words every day.",
    "V19": "Parent app: see every chat, set volume, brightness and sleep, recaps, gallery, custom cards.",
    "V20": "Onboarding 1: meet Cheeko, a tour of the device.",
    "V21": "Onboarding 2: charge, switch on, and Cheeko's own tutorial.",
    "V22": "Onboarding 3: connect with the app, Bluetooth, Wi-Fi and the spoken code.",
    "V23": "Onboarding 4: what to try in the first five minutes.",
    "V25.1": "App guide 1: Home. Today's play, chats word for word, quiz, streak.",
    "V25.2": "App guide 2: Device. Battery, volume, brightness, theme, sounds, sleep mode, Save.",
    "V25.3": "App guide 3: Analytics. Day and week, quiz answers, 12 weeks back.",
    "V25.4": "App guide 4: Gallery and custom cards. Save and share drawings, record your voice on a card.",
    "V25.5": "App guide 5: Profile. Child details, house rules, big days, notifications, help.",
    "V25": "The whole app guide in one film, with YouTube chapters.",
    "X01": "Rejected: too slow for Reels. Kept for reference only.",
}
COMING = [
    ("V24 Onboarding full film", "Planned", "Stitch V20 to V23 into one film, vertical and 16:9, after the final re-render with one music theme (stitch_series.py)."),
    ("V20 to V23 in Hindi", "Planned", "Hindi versions of the four onboarding parts."),
    ("V25 App Guide final pass", "Waiting on the app team", "Re-render the app screens once the Device settings switches and the app's em dashes are fixed, then re-render the parts and re-stitch the films."),
    ("Bachpan ko bachpan rehne do", "In progress", "55 s 16:9 English sticker film. Draft 1 is rendered; not published or numbered yet."),
]


def drive_id(p):
    r = subprocess.run(["xattr", "-p", "com.google.drivefs.item-id#S", str(p)], capture_output=True, text=True)
    return r.stdout.strip() or None


def link(p, folder=False):
    i = drive_id(p) if p and Path(p).exists() else None
    if not i: return None
    return f"https://drive.google.com/drive/folders/{i}" if folder else f"https://drive.google.com/file/d/{i}/view"


def series_of(vid):
    b = vid.replace("-HI", "")
    for name, sr in P.SERIES.items():
        if b in sr["ids"]: return name
    if b.startswith("X"): return "Archive"
    n = int(b[1:].split(".")[0])
    return "Brand and ads" if n <= 10 else "One feature per video" if n <= 18 else "Parent app"


def mmss(s): return f"{int(s // 60)}:{int(round(s % 60)):02d}"


def collect():
    """Every video (plus the archive) with its Drive links and caption, sorted by series."""
    D = P.DRIVE; rows = []
    for v in P.VIDEOS + [P.ARCHIVE]:
        vid, title, slug, repo_dir, fmt, secs, status, voice, notes = v
        root = D / "X Archive" if vid == P.ARCHIVE[0] else P._root_for(vid)
        base = root / f"{vid} {title}"; stem = f"Cheeko_{vid}_{slug}_{fmt}_9x16_{round(secs)}s"; s16 = stem.replace("_9x16_", "_16x9_")
        lang = "Hindi" if vid.endswith("-HI") else "English"
        rp = root / "Ready to post" / lang / f"{vid} {title}"
        cap = P.REPO / repo_dir / "share-copy.txt"
        has16 = (base / "4 Final" / f"{s16}.mp4").exists()
        rows.append(dict(vid=vid, title=title, series=series_of(vid), lang=lang, fmt=fmt, length=mmss(secs), status=status,
                         about=ABOUT.get(vid.replace("-HI", ""), ""),
                         watch=link(base / "4 Final" / f"{stem}.mp4"), w16=link(base / "4 Final" / f"{s16}.mp4") if has16 else None,
                         thumb=link(base / "4 Final" / f"{stem}_thumbnail.jpg"), ready=link(rp, True), files=link(base, True),
                         script=link(base / "1 Script" / f"{vid} {title} - script and shot list.md"),
                         public=link(D / "Public videos" / lang / f"{vid} {title}.mp4"), public16=link(D / "Public videos" / lang / f"{vid} {title} (16x9).mp4"),
                         caption=cap.read_text().strip() if cap.exists() else "",
                         use=("YouTube and website (16:9), WhatsApp and Reels (vertical)" if has16 else
                              "YouTube Shorts, website" if fmt == "Film" else "Instagram Reels, YouTube Shorts, WhatsApp status")))
    for r, v in zip(rows, P.VIDEOS + [P.ARCHIVE]):
        r["repo_dir"] = v[3]
    rows.sort(key=lambda r: (ORDER.index(r["series"]), r["vid"].replace("-HI", ""), r["lang"] == "Hindi"))
    return rows


ORDER = ["Brand and ads", "One feature per video", "Parent app", "Onboarding Cheeko", "Cheeko App Guide", "Archive"]


def build():
    D = P.DRIVE; wb = Workbook(); ws = wb.active; ws.title = "Videos"
    ink, brand, sun, tint = "231A10", "F0521D", "FFC81A", "FFF4E3"
    thin = Side(style="thin", color="EADBC8")
    rows = collect()

    today = datetime.date.today().strftime("%d %b %Y")
    ws["A1"] = "Cheeko videos: master sheet"; ws["A1"].font = Font(name="Arial", size=20, bold=True, color=ink)
    ws["A2"] = (f"Updated {today} · {len(rows)} videos ({sum(r['lang'] == 'English' for r in rows)} English, {sum(r['lang'] == 'Hindi' for r in rows)} Hindi). "
                "Click ▶ Watch to open a video in Drive. Ready to post has only the video, thumbnail and caption. "
                "Links open for anyone the \"cheeko ai videos\" folder is shared with.")
    ws["A2"].font = Font(name="Arial", size=11, color="6B5A48")
    heads = ["ID", "Video", "Series", "Language", "Format", "Length", "Status", "What it's about", "Watch", "16:9 (YouTube)",
             "Thumbnail", "Ready to post", "All files", "Script", "Best for", "Caption (copy and paste)"]
    widths = [9, 30, 20, 10, 8, 8, 16, 58, 11, 15, 12, 14, 11, 10, 34, 70]
    HR = 4
    for c, (h, w) in enumerate(zip(heads, widths), 1):
        cell = ws.cell(HR, c, h); cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=ink)
        cell.alignment = Alignment(vertical="center", wrap_text=True); ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[HR].height = 30
    stat = {"Final draft": ("DDF1DA", "2F6B4F"), "Draft for review": ("FFF1C9", "8A5A00"), "Rejected": ("EEEEEE", "777777")}
    ser = {"Brand and ads": "FFE3D6", "One feature per video": "FFF4E3", "Parent app": "EFE7FF", "Onboarding Cheeko": "E3F2EA", "Cheeko App Guide": "E6F0FF", "Archive": "F2F2F2"}
    for i, r in enumerate(rows, HR + 1):
        vals = [r["vid"], r["title"], r["series"], r["lang"], r["fmt"], r["length"], r["status"], r["about"]]
        for c, val in enumerate(vals, 1):
            cell = ws.cell(i, c, val); cell.font = Font(name="Arial", size=10, bold=(c == 2), color=ink)
            cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 8))
        ws.cell(i, 3).fill = PatternFill("solid", fgColor=ser.get(r["series"], "FFFFFF"))
        bg, fg = stat.get(r["status"], ("FFFFFF", ink)); ws.cell(i, 7).fill = PatternFill("solid", fgColor=bg); ws.cell(i, 7).font = Font(name="Arial", size=10, bold=True, color=fg)
        for c, (key, text) in enumerate([("watch", "▶ Watch"), ("w16", "▶ 16:9"), ("thumb", "Thumbnail"), ("ready", "Ready to post"), ("files", "All files"), ("script", "Script")], 9):
            cell = ws.cell(i, c)
            if r[key]: cell.value = text; cell.hyperlink = r[key]; cell.font = Font(name="Arial", size=10, bold=(key == "watch"), color=brand if key == "watch" else "1A5FB4", underline="single")
            cell.alignment = Alignment(vertical="top")
        ws.cell(i, 15, r["use"]).alignment = Alignment(vertical="top", wrap_text=True); ws.cell(i, 15).font = Font(name="Arial", size=10, color="6B5A48")
        cc = ws.cell(i, 16, r["caption"]); cc.alignment = Alignment(vertical="top", wrap_text=True); cc.font = Font(name="Arial", size=9, color=ink)
        ws.row_dimensions[i].height = 64
        for c in range(1, len(heads) + 1): ws.cell(i, c).border = Border(bottom=thin)
    ws.freeze_panes = ws.cell(HR + 1, 3)
    ws.auto_filter.ref = f"A{HR}:{get_column_letter(len(heads))}{HR + len(rows)}"
    ws.sheet_view.zoomScale = 110

    # Folders and plans
    fs = wb.create_sheet("Folders and plans")
    items = [("Everything (main folder)", "All videos, plans and shared assets", D, True),
             ("Ready to post: English", "Standalone videos: video, thumbnail and caption only", D / "Ready to post/English", True),
             ("Ready to post: Hindi", "Hindi versions of the standalone videos", D / "Ready to post/Hindi", True),
             ("Onboarding Cheeko", "Getting Started series (V20 to V24): plan, every part, its own Ready to post", D / "Onboarding Cheeko", True),
             ("Cheeko App Guide", "Parent app walkthrough (V25): plan, every part, full films, its own Ready to post", D / "Cheeko App Guide", True),
             ("Plans and scripts", "Video ideas, production rules, how we make videos, scripts", D / "00 Plan", True),
             ("Shared assets", "Brand, device photos, real device screens, shipping cards, music", D / "01 Shared assets", True),
             ("Video ideas list", "Every video idea and its status", D / "00 Plan/Cheeko video ideas (all 28).md", False),
             ("Production rules", "Style, facts and the never-do list", D / "00 Plan/Production rules.md", False),
             ("How we make Cheeko videos", "The pipeline, step by step", D / "00 Plan/HOW WE MAKE CHEEKO VIDEOS.md", False),
             ("Onboarding plan and scripts", "V20 to V24", D / "Onboarding Cheeko/00 Plan and scripts", True),
             ("App Guide plan and scripts", "V25, incl. the list of app issues for the app team", D / "Cheeko App Guide/00 Plan and scripts", True),
             ("Archive", "Rejected videos kept for reference", D / "X Archive", True)]
    for c, (h, w) in enumerate(zip(["Folder or document", "What's inside", "Open"], [34, 78, 12]), 1):
        cell = fs.cell(1, c, h); cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=ink); fs.column_dimensions[get_column_letter(c)].width = w
    for i, (name, what, path, folder) in enumerate(items, 2):
        fs.cell(i, 1, name).font = Font(name="Arial", size=11, bold=True, color=ink); fs.cell(i, 2, what).font = Font(name="Arial", size=10, color="6B5A48")
        u = link(path, folder)
        if u: c = fs.cell(i, 3, "Open"); c.hyperlink = u; c.font = Font(name="Arial", size=10, bold=True, color=brand, underline="single")
        fs.row_dimensions[i].height = 22

    # Coming next
    cs = wb.create_sheet("Coming next")
    for c, (h, w) in enumerate(zip(["What", "Status", "Notes"], [34, 28, 100]), 1):
        cell = cs.cell(1, c, h); cell.font = Font(name="Arial", bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor=ink); cs.column_dimensions[get_column_letter(c)].width = w
    for i, row in enumerate(COMING, 2):
        for c, val in enumerate(row, 1):
            cell = cs.cell(i, c, val); cell.font = Font(name="Arial", size=10, bold=(c == 1), color=ink); cell.alignment = Alignment(vertical="top", wrap_text=True)
        cs.row_dimensions[i].height = 34
    for sh in (ws, fs, cs): sh.sheet_properties.tabColor = {"Videos": brand, "Folders and plans": sun, "Coming next": "6C3DFF"}[sh.title]

    out = D / OUT_NAME; wb.save(out)
    missing = [r["vid"] for r in rows if not r["watch"]]
    print("master sheet:", out, len(rows), "videos", ("missing links: " + ", ".join(missing)) if missing else "all linked")
    return out


if __name__ == "__main__":
    build()
